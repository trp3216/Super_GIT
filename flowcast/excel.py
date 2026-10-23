"""สร้างไฟล์ Excel พยากรณ์ที่คำนวณด้วยสูตรในตัว เพิ่มข้อมูลรายวันแล้วผลพยากรณ์อัปเดตเอง"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import BANDS, FEATURE_LABELS, FEATURES, HORIZONS, TARGETS

FONT = "Tahoma"
BLUE = "1F4E79"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "d mmm yyyy"

DATA = "ข้อมูล"
COEF = "ค่าสัมประสิทธิ์"
FC = "พยากรณ์"
ACC = "ความแม่นยำ"

DATA_COLUMNS = [
    ("วันที่", None),
    ("C.2 นครสวรรค์\nระดับน้ำ (ม.รทก.)", "c2_h"),
    ("C.2 นครสวรรค์\nปริมาณน้ำ (ลบ.ม./วิ)", "c2_q"),
    ("เขื่อนเจ้าพระยา\nระดับเหนือเขื่อน (ม.รทก.)", "up_h"),
    ("เขื่อนเจ้าพระยา\nระดับท้ายเขื่อน (ม.รทก.)", "dn_h"),
    ("เขื่อนเจ้าพระยา\nระบาย (ลบ.ม./วิ)", "out_q"),
    ("Ct.19 สะแกกรัง\nระดับน้ำ (ม.รทก.)", "ct19_h"),
    ("Ct.19 สะแกกรัง\nปริมาณน้ำ (ลบ.ม./วิ)", "ct19_q"),
]
DATA_COL = {key: get_column_letter(i + 1) for i, (_, key) in enumerate(DATA_COLUMNS) if key}
TARGET_KEYS = [(t, h) for t in TARGETS for h in HORIZONS]
# แถวในชีตค่าสัมประสิทธิ์
COEF_ROW0 = 4  # ค่าคงที่ แถวถัดไปเป็นตัวแปรตามลำดับ FEATURES
BAND_ROW0 = COEF_ROW0 + len(FEATURES) + 4
RATING_ROW = BAND_ROW0 + len(BANDS) + 3
THRESHOLDS = [("เฝ้าระวัง", 1500, "FFEB9C"), ("เตือนภัย", 2000, "FFC7CE"), ("วิกฤต", 2500, "FF9999")]


def _font(cell, **kw):
    cell.font = Font(name=FONT, **{"size": 10, **kw})


def _header(ws, row, col, text, width=None):
    c = ws.cell(row, col, text)
    _font(c, bold=True, color=BLUE)
    c.fill = HEAD_FILL
    c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def _cell(ws, row, col, value, fmt=None, **font):
    c = ws.cell(row, col, value)
    _font(c, **font)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def _write_data(wb, df):
    ws = wb.create_sheet(DATA)
    for i, (title, _) in enumerate(DATA_COLUMNS, start=1):
        _header(ws, 1, i, title, width=14 if i == 1 else 17)
    ws.row_dimensions[1].height = 44
    ws.freeze_panes = "B2"
    for r, (date, row) in enumerate(df.iterrows(), start=2):
        _cell(ws, r, 1, date.to_pydatetime(), DATE_FMT)
        for i, (_, key) in enumerate(DATA_COLUMNS[1:], start=2):
            v = row[key]
            _cell(ws, r, i, None if pd.isna(v) else float(v), "0.00" if key.endswith("_h") else "#,##0.0")
    return ws


def _coef_col(target, h):
    return 3 + TARGET_KEYS.index((target, h))


def _write_coefficients(wb, model):
    ws = wb.create_sheet(COEF)
    _font(ws.cell(1, 1, "ค่าสัมประสิทธิ์สมการพยากรณ์ (คำนวณจาก python -m flowcast ห้ามแก้ด้วยมือ)"), bold=True, size=12, color=BLUE)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    _header(ws, 3, 1, "ตัวแปร")
    _header(ws, 3, 2, "รหัส")
    for target, h in TARGET_KEYS:
        _header(ws, 3, _coef_col(target, h), f"{TARGETS[target]}\n+{h} วัน", width=16)
    ws.row_dimensions[3].height = 44
    labels = [("ค่าคงที่", "1")] + [
        (f"{FEATURE_LABELS[c]} {'วันนี้' if k == 0 else f'ย้อนหลัง {k} วัน'}", f"{c}[t-{k}]" if k else f"{c}[t]")
        for c, k in FEATURES
    ]
    for i, (label, code) in enumerate(labels):
        _cell(ws, COEF_ROW0 + i, 1, label)
        _cell(ws, COEF_ROW0 + i, 2, code)
        for target, h in TARGET_KEYS:
            _cell(ws, COEF_ROW0 + i, _coef_col(target, h), model["coef"][target][str(h)][i], "0.0000")

    lo_q, hi_q = (int(q * 100) for q in model["quantiles"])
    _font(ws.cell(BAND_ROW0 - 2, 1, f"ช่วงความคลาดเคลื่อน {hi_q - lo_q}% (ค่าที่บวกเพิ่มจากค่าคาดการณ์ แยกตามปริมาณน้ำวันนี้)"), bold=True, color=BLUE)
    _header(ws, BAND_ROW0 - 1, 1, "ปริมาณน้ำวันนี้ตั้งแต่ (ลบ.ม./วิ)")
    _header(ws, BAND_ROW0 - 1, 2, "")
    for target, h in TARGET_KEYS:
        col = 3 + 2 * TARGET_KEYS.index((target, h)) + 6
        _header(ws, BAND_ROW0 - 1, col, f"{TARGETS[target]} +{h}\nต่ำสุด (P{lo_q})", width=16)
        _header(ws, BAND_ROW0 - 1, col + 1, f"{TARGETS[target]} +{h}\nสูงสุด (P{hi_q})", width=16)
    ws.row_dimensions[BAND_ROW0 - 1].height = 44
    for b, start in enumerate(BANDS):
        _cell(ws, BAND_ROW0 + b, 1, start, "#,##0")
        for target, h in TARGET_KEYS:
            col = 3 + 2 * TARGET_KEYS.index((target, h)) + 6
            lo, hi = model["interval"][target][str(h)][b]
            _cell(ws, BAND_ROW0 + b, col, lo, "#,##0")
            _cell(ws, BAND_ROW0 + b, col + 1, hi, "#,##0")

    _font(ws.cell(RATING_ROW - 1, 1, "โค้งระดับน้ำท้ายเขื่อน: ระดับ = a + b·√Q + c·Q (Q = ระบาย ลบ.ม./วิ)"), bold=True, color=BLUE)
    for i, (name, v) in enumerate(zip("abc", model["rating"])):
        _cell(ws, RATING_ROW + i, 1, name)
        _cell(ws, RATING_ROW + i, 2, v, "0.000000")
    return ws


def _band_col(target, h, which):
    return get_column_letter(3 + 2 * TARGET_KEYS.index((target, h)) + 6 + which)


def _write_forecast(wb):
    ws = wb.create_sheet(FC, 0)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJK", [14, 10, 14, 12, 12, 14, 12, 12, 16, 12, 3]):
        ws.column_dimensions[col].width = w
    _font(ws.cell(1, 1, "พยากรณ์ปริมาณน้ำ C.2 นครสวรรค์ และการระบายเขื่อนเจ้าพระยา ล่วงหน้า 1–3 วัน"), bold=True, size=14, color=BLUE)
    _font(ws.cell(2, 1, f"เพิ่มข้อมูลรายวันต่อท้ายชีต \"{DATA}\" แล้วตารางนี้จะคำนวณใหม่จากข้อมูล 4 วันล่าสุดโดยอัตโนมัติ"), italic=True, color="595959")

    # แถวสุดท้ายในชีตข้อมูลที่มีปริมาณน้ำ C.2
    _font(ws.cell(4, 1, "ข้อมูลล่าสุด"), bold=True)
    last = ws.cell(4, 6, f"=MATCH(9.99E+307,'{DATA}'!{DATA_COL['c2_q']}:{DATA_COL['c2_q']})")
    _font(last, color="A6A6A6", size=8)
    _font(ws.cell(4, 5, "แถวที่"), color="A6A6A6", size=8)
    c = ws.cell(4, 2, f"=INDEX('{DATA}'!A:A,$F$4)")
    _font(c, bold=True)
    c.number_format = DATE_FMT
    ws.merge_cells("B4:C4")
    need = [DATA_COL[k] for k in ("c2_q", "out_q", "up_h")]
    counts = "+".join(f"COUNT(INDEX('{DATA}'!{col}:{col},$F$4-3):INDEX('{DATA}'!{col}:{col},$F$4))" for col in need)
    check = ws.cell(
        5, 1,
        f"=IF(AND({counts}={4 * len(need)},INDEX('{DATA}'!A:A,$F$4)-INDEX('{DATA}'!A:A,$F$4-3)=3),"
        f"\"✓ ข้อมูล 4 วันล่าสุดครบ\",\"⚠ ข้อมูล 4 วันล่าสุดไม่ครบหรือวันที่ไม่ต่อเนื่อง ผลพยากรณ์อาจผิด\")",
    )
    _font(check, bold=True, color="C00000")
    ws.merge_cells("A5:J5")

    # ตารางผลพยากรณ์
    heads = ["วันที่", "ล่วงหน้า", "C.2 นครสวรรค์\nคาดการณ์", "ต่ำสุด", "สูงสุด",
             "ระบายท้ายเขื่อน\nคาดการณ์", "ต่ำสุด", "สูงสุด", "ระดับน้ำท้ายเขื่อน\nคาดการณ์ (ม.รทก.)", "สถานะการระบาย"]
    for i, t in enumerate(heads, start=1):
        _header(ws, 7, i, t)
    ws.row_dimensions[7].height = 44
    _font(ws.cell(6, 3, "ปริมาณน้ำ ลบ.ม./วิ  (ต่ำสุด–สูงสุด = ช่วงที่ค่าจริงตกอยู่ 80% ในการทดสอบย้อนหลัง)"), color="595959", size=9)

    feat_row0 = 22  # ค่าคงที่ แถวถัดไปเป็นตัวแปร
    feat_range = f"$B${feat_row0}:$B${feat_row0 + len(FEATURES)}"
    obs_row = {k: 14 + k for k in range(4)}  # 14 = t-3 ... 17 = t
    thr_cells = [f"$M${8 + i}" for i in range(len(THRESHOLDS))]
    for h in HORIZONS:
        r = 7 + h
        _cell(ws, r, 1, f"=$B$4+{h}", DATE_FMT, bold=True)
        _cell(ws, r, 2, f"+{h} วัน").alignment = Alignment(horizontal="center")
        for target, col in (("c2_q", 3), ("out_q", 6)):
            coef_col = get_column_letter(_coef_col(target, h))
            coef = f"'{COEF}'!${coef_col}${COEF_ROW0}:${coef_col}${COEF_ROW0 + len(FEATURES)}"
            pred = _cell(ws, r, col, f"=MAX(0,SUMPRODUCT({feat_range},{coef}))", "#,##0", bold=True, size=11)
            pred.fill = PatternFill("solid", fgColor="E2EFDA")
            cur = "$B$17" if target == "c2_q" else "$E$17"
            band = f"MATCH({cur},'{COEF}'!$A${BAND_ROW0}:$A${BAND_ROW0 + len(BANDS) - 1},1)"
            for j in (0, 1):
                bc = _band_col(target, h, j)
                rng = f"'{COEF}'!${bc}${BAND_ROW0}:${bc}${BAND_ROW0 + len(BANDS) - 1}"
                _cell(ws, r, col + 1 + j, f"=MAX(0,{get_column_letter(col)}{r}+INDEX({rng},{band}))", "#,##0", color="595959")
        a, b, c_ = (f"'{COEF}'!$B${RATING_ROW + i}" for i in range(3))
        _cell(ws, r, 9, f"={a}+{b}*SQRT(F{r})+{c_}*F{r}", "0.00", bold=True)
        status = "\"ปกติ\""
        for (name, _, _), cell in zip(THRESHOLDS, thr_cells):
            status = f"IF(F{r}>={cell},\"{name}\",{status})"
        _cell(ws, r, 10, f"={status}", bold=True).alignment = Alignment(horizontal="center")
    for (name, _, color) in THRESHOLDS:
        ws.conditional_formatting.add("J8:J10", FormulaRule(formula=[f'J8="{name}"'], fill=PatternFill("solid", fgColor=color)))

    # เกณฑ์สถานะ (ผู้ใช้แก้ได้)
    ws.column_dimensions["L"].width = 14
    ws.column_dimensions["M"].width = 12
    _header(ws, 7, 12, "สถานะ")
    _header(ws, 7, 13, "ระบายตั้งแต่\n(ลบ.ม./วิ)")
    for i, (name, value, color) in enumerate(THRESHOLDS):
        _cell(ws, 8 + i, 12, name).fill = PatternFill("solid", fgColor=color)
        _cell(ws, 8 + i, 13, value, "#,##0").fill = INPUT_FILL
    _font(ws.cell(11, 12, "ค่าตั้งต้นตัวอย่าง ปรับตามเกณฑ์\nของหน่วยงานได้ (ช่องสีเหลือง)"), size=8, color="595959")
    ws.cell(11, 12).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("L11:M12")

    # ข้อมูลจริง 4 วันล่าสุด
    _font(ws.cell(12, 1, "ข้อมูลจริง 4 วันล่าสุด (ดึงจากชีตข้อมูล)"), bold=True, color=BLUE)
    obs_cols = [("วันที่", "A", DATE_FMT), ("C.2 นครสวรรค์\n(ลบ.ม./วิ)", DATA_COL["c2_q"], "#,##0"),
                ("เหนือเขื่อน\n(ม.รทก.)", DATA_COL["up_h"], "0.00"), ("ท้ายเขื่อน\n(ม.รทก.)", DATA_COL["dn_h"], "0.00"),
                ("ระบาย\n(ลบ.ม./วิ)", DATA_COL["out_q"], "#,##0"), ("Ct.19 สะแกกรัง\n(ลบ.ม./วิ)", DATA_COL["ct19_q"], "#,##0.0")]
    for i, (t, _, _) in enumerate(obs_cols, start=1):
        _header(ws, 13, i, t)
    ws.row_dimensions[13].height = 44
    for k, r in obs_row.items():
        for i, (_, col, fmt) in enumerate(obs_cols, start=1):
            _cell(ws, r, i, f"=INDEX('{DATA}'!{col}:{col},$F$4-{3 - k})", fmt)
    obs_ref = {"c2_q": "B", "up_h": "C", "out_q": "E", "ct19_q": "F"}

    # ตัวแปรที่ป้อนเข้าสมการ
    _font(ws.cell(feat_row0 - 2, 1, "ค่าที่ป้อนเข้าสมการ (คูณกับค่าสัมประสิทธิ์แล้วรวมกัน)"), bold=True, color=BLUE)
    _header(ws, feat_row0 - 1, 1, "ตัวแปร")
    _header(ws, feat_row0 - 1, 2, "ค่า")
    _cell(ws, feat_row0, 1, "ค่าคงที่")
    _cell(ws, feat_row0, 2, 1)
    for i, (col, lag) in enumerate(FEATURES, start=1):
        _cell(ws, feat_row0 + i, 1, f"{col}[t-{lag}]" if lag else f"{col}[t]", color="595959")
        _cell(ws, feat_row0 + i, 2, f"=N({obs_ref[col]}{obs_row[3 - lag]})", "#,##0.00", color="595959")
    ws.print_area = "A1:M37"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws, obs_row


def _write_chart_table(ws):
    """ตาราง 30 วันล่าสุด (คอลัมน์ O–Q) และพยากรณ์ 3 วัน (คอลัมน์ S–U) สำหรับกราฟ"""
    days = 30
    _font(ws.cell(20, 15, "ข้อมูลสำหรับกราฟ (คำนวณอัตโนมัติ)"), size=8, color="808080")
    for i, t in enumerate(["วันที่", "C.2 จริง", "ระบายจริง", "", "วันที่", "C.2 พยากรณ์", "ระบายพยากรณ์"]):
        if t:
            _header(ws, 21, 15 + i, t, width=12)
    for k in range(days):
        back = days - 1 - k
        for i, col in enumerate(["A", DATA_COL["c2_q"], DATA_COL["out_q"]]):
            _cell(ws, 22 + k, 15 + i, f"=INDEX('{DATA}'!{col}:{col},$F$4-{back})",
                  DATE_FMT if i == 0 else "#,##0", size=8, color="808080")
    # จุดแรกของเส้นพยากรณ์คือข้อมูลจริงวันล่าสุด
    for h in range(4):
        r = 22 + h
        row = ["=$B$4", "=$B$17", "=$E$17"] if h == 0 else [f"=$B$4+{h}", f"=C{7 + h}", f"=F{7 + h}"]
        for i, v in enumerate(row):
            _cell(ws, r, 19 + i, v, DATE_FMT if i == 0 else "#,##0", size=8, color="808080")

    chart = ScatterChart()
    chart.title = "30 วันล่าสุด และพยากรณ์ 3 วัน (ลบ.ม./วิ)"
    chart.style = 13
    chart.height, chart.width = 9, 22
    chart.y_axis.title = "ลบ.ม./วิ"
    chart.y_axis.majorGridlines = None
    chart.x_axis.number_format = "d/m"
    chart.x_axis.majorUnit = 7
    chart.x_axis.majorGridlines = None
    for x_col, y_col, last, title, color, dash in (
        (15, 16, 21 + days, "C.2 จริง", "2E75B6", None),
        (15, 17, 21 + days, "ระบายจริง", "C55A11", None),
        (19, 20, 25, "C.2 พยากรณ์", "2E75B6", "dash"),
        (19, 21, 25, "ระบายพยากรณ์", "C55A11", "dash"),
    ):
        xs = Reference(ws, min_col=x_col, min_row=22, max_row=last)
        ys = Reference(ws, min_col=y_col, min_row=22, max_row=last)
        series = Series(ys, xs, title=title)
        series.marker.symbol = "none"
        series.smooth = False
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 22000
        if dash:
            series.graphicalProperties.line.dashStyle = dash
        chart.series.append(series)
    ws.add_chart(chart, "D19")


def _write_accuracy(wb, model, backtests):
    ws = wb.create_sheet(ACC)
    _font(ws.cell(1, 1, "ความแม่นยำจากการทดสอบย้อนหลังแบบเว้นทีละปี"), bold=True, size=12, color=BLUE)
    _font(ws.cell(2, 1, f"ฝึกแบบจำลองด้วยปีอื่นทั้งหมด แล้วพยากรณ์ปีที่เว้นไว้ วนครบทุกปี (ข้อมูล {model['trained_on'][0]} ถึง {model['trained_on'][1]})"), color="595959")
    _font(ws.cell(3, 1, "\"ไม่ใช้แบบจำลอง\" = เดาว่าค่าในอนาคตเท่ากับวันนี้ ใช้เป็นเกณฑ์เทียบ"), color="595959")
    heads = ["พยากรณ์", "ล่วงหน้า (วัน)", "คลาดเฉลี่ย\n(ลบ.ม./วิ)", "ไม่ใช้แบบจำลอง\nคลาดเฉลี่ย", "ดีขึ้น", "คลาดเฉลี่ยเมื่อน้ำ\n> 1,000 ลบ.ม./วิ", "ไม่ใช้แบบจำลอง\n> 1,000", "จำนวนวัน\nที่ทดสอบ"]
    widths = [26, 12, 14, 16, 10, 18, 16, 12]
    for i, (t, w) in enumerate(zip(heads, widths), start=1):
        _header(ws, 5, i, t, width=w)
    ws.row_dimensions[5].height = 44
    for r, m in enumerate(model["metrics"], start=6):
        vals = [TARGETS[m["target"]], m["h"], m["mae"], m["mae_persist"], 1 - m["mae"] / m["mae_persist"],
                m["mae_high"], m["mae_persist_high"], m["n"]]
        fmts = [None, "0", "#,##0.0", "#,##0.0", "0%", "#,##0.0", "#,##0.0", "#,##0"]
        for i, (v, f) in enumerate(zip(vals, fmts), start=1):
            _cell(ws, r, i, v, f)

    # ผลทดสอบย้อนหลังรายวัน (เรียงตามวันที่เป้าหมาย)
    series = {}
    for (target, h), bt in backtests.items():
        if h in (1, 3):
            s = bt["pred"].copy()
            s.index = s.index + pd.Timedelta(days=h)
            series[f"{TARGETS[target]} พยากรณ์ +{h} วัน"] = s
        if h == 1:
            obs = bt["obs"].copy()
            obs.index = obs.index + pd.Timedelta(days=1)
            series[f"{TARGETS[target]} จริง"] = obs
    order = [f"{TARGETS[t]} {s}" for t in ("out_q", "c2_q") for s in ("จริง", "พยากรณ์ +1 วัน", "พยากรณ์ +3 วัน")]
    table = pd.DataFrame(series)[order].sort_index()
    row0 = 16
    _font(ws.cell(row0 - 1, 1, "ผลทดสอบย้อนหลังรายวัน (ค่าพยากรณ์วางตรงกับวันที่ที่พยากรณ์ถึง)"), bold=True, color=BLUE)
    _header(ws, row0, 10, "วันที่", width=12)
    for i, name in enumerate(order, start=11):
        _header(ws, row0, i, name, width=14)
    ws.row_dimensions[row0].height = 44
    for r, (date, row) in enumerate(table.iterrows(), start=row0 + 1):
        ws.cell(r, 10, date.to_pydatetime()).number_format = DATE_FMT
        for i, v in enumerate(row.values, start=11):
            ws.cell(r, i, None if np.isnan(v) else round(float(v), 1)).number_format = "#,##0"

    # กราฟฤดูน้ำหลากปีล่าสุดที่มีข้อมูลครบ
    flood = table.index[(table.index.month >= 8) & (table.index.month <= 12)]
    year = max(y for y in set(flood.year) if (flood.year == y).sum() >= 140)
    rows = [row0 + 1 + i for i, d in enumerate(table.index) if d.year == year and 8 <= d.month <= 12]
    chart = LineChart()
    chart.title = f"ระบายท้ายเขื่อนเจ้าพระยา ส.ค.–ธ.ค. {year + 543}: จริง vs พยากรณ์"
    chart.height, chart.width = 9, 24
    chart.y_axis.title = "ลบ.ม./วิ"
    chart.y_axis.majorGridlines = None
    for col, name in zip((11, 12, 13), order):
        chart.series.append(Series(Reference(ws, min_col=col, min_row=rows[0], max_row=rows[-1]), title=name))
    chart.set_categories(Reference(ws, min_col=10, min_row=rows[0], max_row=rows[-1]))
    chart.x_axis.number_format = "d/m"
    for s, color, dash in zip(chart.series, ("404040", "2E75B6", "C55A11"), (None, "dash", "sysDot")):
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 19000
        if dash:
            s.graphicalProperties.line.dashStyle = dash
    ws.add_chart(chart, f"A{row0}")
    ws.freeze_panes = f"A{row0 + 1}"


def _write_guide(wb, model):
    ws = wb.create_sheet("วิธีใช้")
    ws.column_dimensions["A"].width = 110
    lines = [
        ("วิธีใช้", True),
        (f"1. เพิ่มข้อมูลรายวันต่อท้ายชีต \"{DATA}\" ให้ครบอย่างน้อยคอลัมน์ วันที่, C.2 ปริมาณน้ำ, ระดับเหนือเขื่อน, ระบาย และ Ct.19 ปริมาณน้ำ (ว่างได้ถ้าไม่มี)", False),
        (f"2. เปิดชีต \"{FC}\" ผลพยากรณ์ 1–3 วันถัดจากวันล่าสุดคำนวณให้อัตโนมัติ", False),
        ("3. ถ้ามีข้อความเตือนสีแดงว่าข้อมูลไม่ครบ ให้ตรวจว่า 4 วันล่าสุดไม่มีช่องว่างและวันที่เรียงต่อเนื่อง", False),
        ("", False),
        ("หลักการ", True),
        ("• สมการเชิงเส้น: ค่าพยากรณ์ = ค่าคงที่ + Σ(ค่าสัมประสิทธิ์ × ข้อมูล) ใช้ข้อมูลวันนี้ เมื่อวาน และ 3 วันก่อน", False),
        ("• ตัวแปร: ปริมาณน้ำ C.2 นครสวรรค์, การระบายท้ายเขื่อน, ระดับน้ำเหนือเขื่อน, ปริมาณน้ำ Ct.19 สะแกกรัง", False),
        (f"• แยกสมการตามจุดพยากรณ์และจำนวนวันล่วงหน้า รวม 6 สมการ ดูค่าได้ที่ชีต \"{COEF}\"", False),
        ("• ช่วงต่ำสุด–สูงสุด มาจากความคลาดเคลื่อนจริงในการทดสอบย้อนหลัง (เปอร์เซ็นไทล์ 10 และ 90) แยกตามปริมาณน้ำวันนี้", False),
        ("• ระดับน้ำท้ายเขื่อนคาดการณ์ คำนวณจากค่าระบายคาดการณ์ด้วยโค้งความสัมพันธ์ระดับน้ำ–ปริมาณน้ำ", False),
        ("", False),
        ("ข้อจำกัด", True),
        ("• การระบายเป็นการตัดสินใจของผู้ปฏิบัติงาน แบบจำลองเรียนรู้จากการสั่งการในอดีต ถ้าเกณฑ์บริหารน้ำเปลี่ยน ความแม่นยำจะลดลง", False),
        ("• ไม่มีข้อมูลต้นน้ำ (เขื่อนภูมิพล สิริกิติ์ สถานีปิง วัง ยม น่าน ฝน) จึงพยากรณ์ไกลเกิน 3 วันได้ไม่ดี และจะตามไม่ทันเมื่อน้ำเพิ่มเร็ว", False),
        (f"• ฝึกด้วยข้อมูล {model['trained_on'][0]} ถึง {model['trained_on'][1]} เมื่อมีข้อมูลใหม่มากพอ ให้ฝึกใหม่ด้วยคำสั่ง python -m flowcast --excel <ไฟล์ข้อมูล>", False),
        ("• ใช้ประกอบการตัดสินใจเท่านั้น ควรตรวจสอบกับประกาศและการคาดการณ์ทางการของกรมชลประทาน", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(r, 1, text)
        _font(c, bold=bold, size=12 if bold else 10, color=BLUE if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(df, model, backtests, path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_data(wb, df.loc[model["trained_on"][0]:model["trained_on"][1]])
    _write_coefficients(wb, model)
    ws, _ = _write_forecast(wb)
    _write_chart_table(ws)
    _write_accuracy(wb, model, backtests)
    _write_guide(wb, model)
    wb.active = 0
    wb.save(path)
