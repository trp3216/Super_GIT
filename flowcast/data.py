"""อ่านและทำความสะอาดไฟล์ Excel ฐานข้อมูลจัดสรรน้ำ"""
import datetime as dt

import numpy as np
import openpyxl
import pandas as pd

COLUMNS = ["c2_h", "c2_q", "up_h", "dn_h", "out_q", "ct19_h", "ct19_q"]

# ช่วงค่าที่เป็นไปได้ นอกช่วงนี้ถือว่าพิมพ์ผิด (เช่น 1645 แทน 16.45)
VALID_RANGE = {
    "c2_h": (10, 30),
    "up_h": (8, 25),
    "dn_h": (2, 20),
    "c2_q": (0, 6000),
    "out_q": (0, 6000),
    "ct19_q": (0, 2000),
}


def _read_legacy_sheet(ws):
    """ชีตปี 2540–2562: มีเฉพาะข้อมูลเขื่อนเจ้าพระยา วันที่บางแถวเป็นข้อความหรือปีผิด

    ข้อมูลเป็นรายวันต่อเนื่อง จึงใช้วันที่จากเซลล์เมื่อสอดคล้องกับแถวก่อนหน้า
    ไม่เช่นนั้นใช้วันถัดจากแถวก่อนหน้า
    """
    rows, prev = [], None
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[3] is None and r[5] is None:
            continue
        expected = prev + dt.timedelta(days=1) if prev else dt.datetime(1997, 1, 1)
        date = r[2] if isinstance(r[2], dt.datetime) else r[1]
        if not (isinstance(date, dt.datetime) and abs((date - expected).days) <= 3):
            date = expected
        prev = date
        rows.append((date, None, None, r[3], r[4], r[5], None, None))
    return rows


def _read_year_sheet(ws):
    """ชีตรายปีตั้งแต่ปี 2563: วันที่, C.2 ระดับ/ปริมาณ, C.13 เหนือ/ท้าย/ระบาย, Ct.19 ระดับ/ปริมาณ"""
    return [
        tuple(r[:8])
        for r in ws.iter_rows(min_row=4, values_only=True)
        if isinstance(r[0], dt.datetime) and r[5] is not None
    ]


def clean(df):
    """ตัดค่าที่พิมพ์ผิดให้เป็นค่าว่าง แล้วจัดเป็นข้อมูลรายวันต่อเนื่อง"""
    df = df.apply(pd.to_numeric, errors="coerce")
    for col, (lo, hi) in VALID_RANGE.items():
        df.loc[(df[col] < lo) | (df[col] > hi), col] = np.nan
    # ค่าระบายที่กระโดดขึ้นวันเดียว เช่น 4338 ระหว่างวันที่ระบาย ~200
    q = df["out_q"]
    neighbours = pd.concat([q.shift(k) for k in (-3, -2, -1, 1, 2, 3)], axis=1).median(axis=1)
    df.loc[(q > 500) & (q > 5 * neighbours), "out_q"] = np.nan
    return df.asfreq("D")


def load_excel(path):
    """คืน DataFrame รายวัน index เป็นวันที่ คอลัมน์ตาม COLUMNS"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        header = next(ws.iter_rows(max_row=1, values_only=True))
        rows += _read_year_sheet(ws) if header[0] == "วันที่" else _read_legacy_sheet(ws)
    df = pd.DataFrame(rows, columns=["date"] + COLUMNS)
    df["date"] = pd.to_datetime(df["date"]).dt.normalize()
    df = df.drop_duplicates("date", keep="last").set_index("date").sort_index()
    return clean(df)
