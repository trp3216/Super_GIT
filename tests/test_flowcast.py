import datetime as dt

import numpy as np
import openpyxl
import pandas as pd

from flowcast.data import clean, load_excel
from flowcast.excel import build_workbook
from flowcast.model import band_index, features, forecast, train


def synthetic(days=4 * 365, seed=0):
    """C.2 เป็นคลื่นตามฤดูกาล การระบายตามน้ำ C.2 เมื่อวานลบน้ำที่ผันเข้าคลอง"""
    rng = np.random.default_rng(seed)
    idx = pd.date_range("2020-01-01", periods=days, freq="D")
    season = np.clip(np.sin(2 * np.pi * (idx.dayofyear - 180) / 365), 0, None)
    c2 = 200 + 2000 * season + rng.normal(0, 20, days).cumsum() * 0.1
    out = np.clip(np.r_[c2[0], c2[:-1]] - 150, 0, None)
    return pd.DataFrame({
        "c2_h": 16 + c2 / 500, "c2_q": c2, "up_h": 16 + rng.normal(0, 0.1, days),
        "dn_h": 3.6 + 0.21 * np.sqrt(out), "out_q": out, "ct19_h": np.nan, "ct19_q": np.nan,
    }, index=idx)


def test_features_use_past_values_only():
    df = synthetic(10)
    X = features(df)
    t = df.index[5]
    assert X.loc[t, "c2_q[t]"] == df.loc[t, "c2_q"]
    assert X.loc[t, "c2_q[t-3]"] == df["c2_q"].iloc[2]
    assert X.loc[t, "ct19_q[t]"] == 0  # ค่าว่างของ Ct.19 ถือเป็น 0


def test_clean_removes_typos_and_spikes():
    idx = pd.date_range("2001-07-10", periods=9, freq="D")
    df = pd.DataFrame({c: np.nan for c in ["c2_h", "c2_q", "dn_h", "ct19_h", "ct19_q"]}, index=idx)
    df["up_h"] = [16.5, 16.4, 1645, 16.4, 16.4, 16.3, 16.3, 16.2, 16.2]
    df["out_q"] = [100, 120, 200, 250, 300, 4338, 195, 84, 68]
    out = clean(df)
    assert np.isnan(out.loc["2001-07-12", "up_h"])
    assert np.isnan(out.loc["2001-07-15", "out_q"])
    assert out.loc["2001-07-14", "out_q"] == 300


def test_band_index():
    assert [band_index(v) for v in (0, 499, 500, 1499, 1500, 4000)] == [0, 0, 1, 1, 2, 2]


def test_train_beats_persistence_and_forecasts():
    df = synthetic()
    model, backtests = train(df)
    for m in model["metrics"]:
        assert m["mae"] < m["mae_persist"]
    t, fc = forecast(model, df)
    assert t == df.index[-1]
    assert list(fc.index) == [t + pd.Timedelta(days=h) for h in (1, 2, 3)]
    # การระบายพรุ่งนี้ = C.2 วันนี้ − 150 ตามที่สร้างข้อมูลไว้
    assert abs(fc["out_q"].iloc[0] - (df["c2_q"].iloc[-1] - 150)) < 5
    assert (fc["out_q_low"] <= fc["out_q"]).all() and (fc["out_q"] <= fc["out_q_high"]).all()


def test_load_excel_repairs_legacy_dates(tmp_path):
    wb = openpyxl.Workbook()
    legacy = wb.active
    for _ in range(4):
        legacy.append([None] * 6)
    legacy.append([None, dt.datetime(1997, 1, 1), dt.datetime(1997, 1, 1), 16.0, 6.0, 50])
    legacy.append([None, None, "  2 ม.ค. 40", 16.1, 6.1, 52])
    legacy.append([None, None, dt.datetime(1900, 1, 3), 16.2, 6.2, 54])
    year = wb.create_sheet("ปี63")
    year.append(["วันที่", "C.2 นครสวรรค์", None, "C.13 เขื่อนเจ้าพระยา"])
    year.append([None])
    year.append([None])
    year.append([dt.datetime(2020, 1, 1), 17.5, 257, 14.06, 5.78, 90, 0, 0])
    year.append([dt.datetime(2020, 12, 31), None, None, None, None, None])
    path = tmp_path / "in.xlsx"
    wb.save(path)

    df = load_excel(path)
    assert df.loc["1997-01-02", "out_q"] == 52
    assert df.loc["1997-01-03", "out_q"] == 54
    assert df.loc["2020-01-01", "c2_q"] == 257
    assert df.index[-1] == pd.Timestamp("2020-01-01")


def test_build_workbook(tmp_path):
    df = synthetic()
    model, backtests = train(df)
    path = tmp_path / "fc.xlsx"
    build_workbook(df, model, backtests, path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[0] == "พยากรณ์"
    ws = wb["พยากรณ์"]
    assert ws["C8"].value.startswith("=MAX(0,SUMPRODUCT($B$22:$B$32,'ค่าสัมประสิทธิ์'!$C$4:$C$14")
    assert ws["F10"].value.startswith("=MAX(0,SUMPRODUCT($B$22:$B$32,'ค่าสัมประสิทธิ์'!$H$4:$H$14")
    assert wb["ข้อมูล"].max_row == len(df) + 1
